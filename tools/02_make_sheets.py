"""
2단계: 판정자별 판정 시트(엑셀)를 만든다.

사용법:
    pip install pandas openpyxl
    python 02_make_sheets.py A v1      # 세트 A, 가이드 v1 → sheet_A_v1_R1.xlsx ~ R4
    python 02_make_sheets.py B v2      # 세트 B, 가이드 v2
    python 02_make_sheets.py A v2      # 세트 A 재판정 (2주 뒤)

- 판정자마다 영상 순서를 다르게 섞는다 (순서 효과 방지).
- 화질 점수는 넣지 않는다 (판정자가 보면 안 됨).
- 칸마다 드롭다운이 걸려 있어 오타가 나지 않는다.
"""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RATERS = ["R1", "R2", "R3", "R4"]
SEED = 2026

COLUMNS = [
    ("순번", 6), ("video_id", 10), ("영상 링크", 40),
    ("왜곡 유무", 12), ("신체 형태", 10), ("신체 움직임", 11),
    ("사물·배경", 10), ("화질", 8), ("심각도", 10),
    ("재생 횟수", 10), ("메모", 50),
]


def make_sheet(videos: pd.DataFrame, set_name: str, guide: str, rater: str, order_seed: int):
    order = videos.sample(frac=1, random_state=order_seed).reset_index(drop=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "판정"

    head_fill = PatternFill("solid", fgColor="E7F1FB")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "D2"

    for i, row in order.iterrows():
        r = i + 2
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=row["video_id"])
        link = ws.cell(row=r, column=3, value=row["video_link"])
        link.hyperlink = row["video_link"]
        link.font = Font(color="0563C1", underline="single")

    last = len(order) + 1
    dv_presence = DataValidation(type="list", formula1='"없음,있음,판단 불가"', allow_blank=True)
    dv_check = DataValidation(type="list", formula1='"O"', allow_blank=True)
    dv_sev = DataValidation(type="list", formula1='"경미,심각"', allow_blank=True)
    dv_count = DataValidation(type="whole", operator="between", formula1="1", formula2="20",
                              allow_blank=True)
    for dv in (dv_presence, dv_check, dv_sev, dv_count):
        ws.add_data_validation(dv)
    dv_presence.add(f"D2:D{last}")
    dv_check.add(f"E2:H{last}")
    dv_sev.add(f"I2:I{last}")
    dv_count.add(f"J2:J{last}")

    info = wb.create_sheet("안내")
    lines = [
        f"세트 {set_name} · 가이드 {guide} · 판정자 {rater}",
        "",
        "1. 영상 링크를 열어 필요한 만큼 보세요. 본 횟수를 '재생 횟수'에 적습니다.",
        "2. 현실에서는 일어날 수 없는 모양이나 움직임이 보이면 '있음'입니다.",
        "3. '있음'이면 해당하는 유형 칸에 모두 O를 고릅니다 (여러 개 가능).",
        "4. 가장 눈에 띄는 왜곡 기준으로 심각도를 고릅니다.",
        "5. 애매했다면 이유를 '메모'에 한 줄 적어 주세요. 가장 중요한 자료입니다.",
        "6. '없음'이면 유형과 심각도는 비워 둡니다.",
        "7. 다른 판정자와 상의하지 마세요.",
    ]
    for i, text in enumerate(lines, start=1):
        info.cell(row=i, column=1, value=text)
    info.column_dimensions["A"].width = 90

    out = f"sheet_{set_name}_{guide}_{rater}.xlsx"
    wb.save(out)
    return out


def main():
    if len(sys.argv) != 3:
        print("사용법: python 02_make_sheets.py <A|B> <v1|v2>")
        sys.exit(1)
    set_name, guide = sys.argv[1].upper(), sys.argv[2].lower()
    videos = pd.read_csv(f"videos_{set_name}.csv")
    for k, rater in enumerate(RATERS):
        seed = SEED + k * 101 + (0 if guide == "v1" else 7) + (0 if set_name == "A" else 13)
        print("만듦:", make_sheet(videos, set_name, guide, rater, seed))


if __name__ == "__main__":
    main()
